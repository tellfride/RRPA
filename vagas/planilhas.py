"""Planilhas (.xlsx): modelo de cadastro de cada cargo e exportação dos colaboradores da vaga."""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.utils import formatar_cpf, formatar_telefone, texto_de_planilha, texto_prazo

# (título da coluna, obrigatória, largura, dica)
COLUNAS_MODELO = [
    ("Nome completo", True, 34, "Nome e sobrenome."),
    ("CPF", True, 17, "Com ou sem pontos e traço."),
    ("Telefone (WhatsApp)", True, 21, "Com DDD. Ex.: (11) 91234-5678"),
    ("CEP", True, 12, "Com ou sem traço. O endereço é buscado a partir dele."),
    ("Endereço (rua ou avenida)", False, 36, "Pode deixar em branco: o CEP preenche."),
    ("Número", False, 10, ""),
    ("Complemento", False, 18, ""),
    ("Bairro", False, 22, "Pode deixar em branco: o CEP preenche."),
    ("Cidade", False, 22, "Pode deixar em branco: o CEP preenche."),
    ("UF", False, 6, "Sigla do estado. Ex.: SP"),
]
UFS = "AC,AL,AP,AM,BA,CE,DF,ES,GO,MA,MT,MS,MG,PA,PB,PR,PE,PI,RJ,RN,RS,RO,RR,SC,SP,SE,TO".split(",")
LINHAS_FORMATADAS = 500
VERDE = "0E6B66"


def _cabecalho(ws, colunas=COLUNAS_MODELO, linha=1):
    borda = Side(style="thin", color="C3CDC9")
    for i, (titulo, obrigatoria, largura, dica) in enumerate(colunas, start=1):
        celula = ws.cell(row=linha, column=i, value=f"{titulo} *" if obrigatoria else titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=VERDE)
        celula.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        celula.border = Border(bottom=borda)
        if dica:
            celula.comment = Comment(dica, "R-RPA")
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha].height = 32


def gerar_modelo(vaga, config):
    """Planilha de cadastro do cargo: Candidatos (para preencher), Instruções e uma aba oculta que identifica a vaga."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidatos"
    _cabecalho(ws)
    ws.freeze_panes = "A2"
    for linha in range(2, LINHAS_FORMATADAS + 2):  # tudo como texto: o Excel não come zeros de CPF e CEP
        for coluna in range(1, len(COLUNAS_MODELO) + 1):
            ws.cell(row=linha, column=coluna).number_format = "@"
    validacao = DataValidation(type="list", formula1='"' + ",".join(UFS) + '"', allow_blank=True)
    validacao.error = "Use a sigla do estado (ex.: SP)."
    ws.add_data_validation(validacao)
    validacao.add(f"J2:J{LINHAS_FORMATADAS + 1}")

    ajuda = wb.create_sheet("Instruções")
    ajuda.column_dimensions["A"].width = 110
    titulo = ajuda.cell(row=1, column=1, value=f"Modelo de cadastro — {vaga.nome}")
    titulo.font = Font(bold=True, size=15, color=VERDE)
    linhas = [
        "",
        "Como preencher",
        "1. Use a aba “Candidatos”, a partir da linha 2. Uma pessoa por linha.",
        "2. Campos com * são obrigatórios: nome completo, CPF, telefone com DDD e CEP.",
        "3. CPF e CEP podem ter pontos e traço ou só números. Não altere os títulos nem a ordem das colunas.",
        "4. Deixe rua, bairro, cidade e UF em branco se quiser: o sistema busca pelo CEP.",
        "5. Cada pessoa será cadastrada como candidata a esta vaga. CPFs que já existem no sistema não são duplicados.",
        f"6. Quem for chamado poderá ficar até {config.prazo_texto} na vaga, a contar do mês da chamada.",
        f"7. Depois de preencher, envie o arquivo em Vagas → {vaga.nome} → Importar planilha.",
        "",
        "Exemplo de linha (fictício)",
    ]
    for i, texto in enumerate(linhas, start=2):
        celula = ajuda.cell(row=i, column=1, value=texto)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        if texto in ("Como preencher", "Exemplo de linha (fictício)"):
            celula.font = Font(bold=True, size=12)
    exemplo = ajuda.cell(
        row=len(linhas) + 2, column=1,
        value="Maria da Silva | 529.982.247-25 | (11) 91234-5678 | 01310-100 | Avenida Paulista | 1000 | Sala 2 | Bela Vista | São Paulo | SP",
    )
    exemplo.font = Font(italic=True, color="44585C")
    if vaga.descricao:
        extra = ajuda.cell(row=len(linhas) + 4, column=1, value=f"Sobre a vaga: {vaga.descricao}")
        extra.alignment = Alignment(wrap_text=True, vertical="top")

    meta = wb.create_sheet("_rrpa")
    meta["A1"], meta["B1"] = "RRPA", "modelo-cadastro"
    meta["A2"], meta["B2"] = "vaga_id", vaga.pk
    meta["A3"], meta["B3"] = "vaga_nome", vaga.nome
    meta.sheet_state = "hidden"

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


def exportar_vaga(vaga, vinculos, incluir_cpf):
    """Colaboradores da vaga (já preparados por vinculos.apresentacao.preparar) em uma planilha."""
    wb = Workbook()
    ws = wb.active
    ws.title = vaga.nome[:31]
    colunas = [("Colaborador", False, 34, ""), ("CPF", False, 17, ""), ("Telefone", False, 18, ""),
               ("Chamado em", False, 14, ""), ("Prazo até", False, 14, ""), ("Situação", False, 16, ""),
               ("Prazo", False, 20, ""), ("Check-ins", False, 46, "")]
    if not incluir_cpf:
        colunas = [c for c in colunas if c[0] != "CPF"]
    _cabecalho(ws, colunas)
    ws.freeze_panes = "A2"
    for v in vinculos:
        checkins = "; ".join(f'{c["rotulo"]}: {c["descricao"].split(" (")[0].lower()}' for c in v.chips)
        linha = [
            v.candidato.nome, formatar_cpf(v.candidato.cpf), formatar_telefone(v.candidato.telefone),
            v.data_chamada.strftime("%d/%m/%Y"), v.data_fim.strftime("%d/%m/%Y"), v.sit_rotulo,
            texto_prazo(v.dias) if v.status == "ativo" else "—", checkins,
        ]
        if not incluir_cpf:
            del linha[1]  # a coluna de CPF só existe para quem pode ver o documento
        ws.append([texto_de_planilha(valor) for valor in linha])
    ws.append([])
    ws.append([f"Exportado em {timezone.localtime():%d/%m/%Y %H:%M}"])
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
