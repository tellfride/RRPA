from datetime import timedelta
from io import BytesIO
from unittest import mock

from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models import ProtectedError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from candidatos import services
from candidatos.models import Candidato, Importacao, LinkCadastro
from contas.models import Usuario
from core.excecoes import RegraDeNegocio
from core.fabrica import cpf_valido, telefone_ficticio
from core.models import Configuracao, LogAtividade
from mensagens.models import Disparo, EnvioMensagem, MensagemDisparo, MensagemRecebida
from vagas import planilhas
from vagas.models import Vaga
from vinculos import services as vinculos
from vinculos.models import Checkin, Vinculo


def dados_publicos(n=1, **extra):
    base = {
        "nome": f"Maria Souza {n}", "cpf": cpf_valido(500 + n), "telefone": telefone_ficticio(n), "email": "",
        "cep": "01310100", "logradouro": "Avenida Paulista", "numero": "1000", "complemento": "", "bairro": "Bela Vista",
        "cidade": "São Paulo", "uf": "SP", "aceite": "on", "website": "",
    }
    base.update(extra)
    return base


class BaseTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        Configuracao.atual()
        cls.rec = Usuario.objects.create_user("rec", password="x", first_name="Marina", last_name="Lopes", perfil="recrutador")
        cls.gestor = Usuario.objects.create_user("ges", password="x", perfil="gestor")
        cls.vaga = Vaga.objects.create(nome="Farmácia", quantidade=3)


class CadastroPorLinkTests(BaseTests):
    def test_link_pessoal_cadastra_e_associa_a_quem_forneceu(self):
        link = LinkCadastro.pessoal_de(self.rec)
        pagina = self.client.get(link.caminho)
        self.assertContains(pagina, "Marina Lopes")  # o candidato vê de quem recebeu o link
        resposta = self.client.post(link.caminho, dados_publicos(1, vaga=self.vaga.pk, aceita_whatsapp="on"))
        self.assertRedirects(resposta, reverse("candidatos:cadastro_concluido"))
        c = Candidato.objects.get()
        self.assertEqual((c.indicado_por, c.link, c.origem, c.status), (self.rec, link, "link_pessoal", "pendente"))
        self.assertTrue(c.whatsapp_ok)
        self.assertIsNotNone(c.consentimento_em)
        self.assertEqual(list(c.vagas_interesse.all()), [self.vaga])
        link.refresh_from_db()
        self.assertIsNone(link.usado_em)  # o link pessoal continua valendo

    def test_whatsapp_fica_como_nao_aceito_quando_nao_marcado(self):
        link = LinkCadastro.pessoal_de(self.rec)
        self.client.post(link.caminho, dados_publicos(1))
        self.assertIs(Candidato.objects.get().whatsapp_ok, False)

    def test_convite_depois_da_aprovacao_entra_aprovado_e_so_funciona_uma_vez(self):
        link = services.criar_convite(self.rec, "João Teixeira", telefone_ficticio(7), vaga=self.vaga)
        self.assertEqual(link.aprovado_por, self.rec)
        self.client.post(link.caminho, dados_publicos(2))
        c = Candidato.objects.get()
        self.assertEqual((c.status, c.origem, c.aprovado_por, c.indicado_por), ("aprovado", "convite", self.rec, self.rec))
        link.refresh_from_db()
        self.assertEqual(link.estado, "usado")
        self.assertEqual(self.client.get(link.caminho).status_code, 410)
        self.assertEqual(self.client.post(link.caminho, dados_publicos(3)).status_code, 410)
        self.assertEqual(Candidato.objects.count(), 1)

    def test_convite_expirado_ou_desativado_nao_abre(self):
        expirado = services.criar_convite(self.rec, "Ana Lima", telefone_ficticio(8))
        LinkCadastro.objects.filter(pk=expirado.pk).update(expira_em=timezone.now() - timedelta(minutes=1))
        self.assertEqual(self.client.get(expirado.caminho).status_code, 410)
        desativado = services.criar_convite(self.rec, "Ana Lima", telefone_ficticio(8))
        services.revogar_link(desativado, self.rec)
        self.assertEqual(self.client.get(desativado.caminho).status_code, 410)
        self.assertEqual(self.client.get("/cadastro/inexistente/").status_code, 404)

    def test_convite_de_candidato_existente_atualiza_o_mesmo_cadastro(self):
        c = Candidato.objects.create(nome="Pedro Alves", cpf=cpf_valido(900), telefone=telefone_ficticio(9), cep="01310100")
        link = services.criar_convite(self.rec, c.nome, c.telefone, candidato=c)
        c.refresh_from_db()
        self.assertEqual(c.status, "aprovado")  # gerar o convite é a aprovação
        self.client.post(link.caminho, dados_publicos(9, nome="Pedro Alves Junior", cpf=c.cpf, numero="55"))
        c.refresh_from_db()
        self.assertEqual((Candidato.objects.count(), c.nome, c.numero, c.indicado_por), (1, "Pedro Alves Junior", "55", self.rec))

    def test_convite_de_candidato_exige_o_mesmo_cpf(self):
        c = Candidato.objects.create(nome="Pedro Alves", cpf=cpf_valido(900), telefone=telefone_ficticio(9), cep="01310100")
        link = services.criar_convite(self.rec, c.nome, c.telefone, candidato=c)
        resposta = self.client.post(link.caminho, dados_publicos(9, cpf=cpf_valido(901)))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "não confere")

    def test_cpf_repetido_nao_sobrescreve_dados_nem_revela_que_existe(self):
        existente = Candidato.objects.create(nome="Maria Original", cpf=cpf_valido(501), telefone=telefone_ficticio(1), cep="01310100")
        link = LinkCadastro.pessoal_de(self.rec)
        resposta = self.client.post(link.caminho, dados_publicos(1, nome="Outra Pessoa", vaga=self.vaga.pk))
        self.assertRedirects(resposta, reverse("candidatos:cadastro_concluido"))  # mesma resposta de um cadastro novo
        existente.refresh_from_db()
        self.assertEqual((Candidato.objects.count(), existente.nome, existente.indicado_por), (1, "Maria Original", self.rec))
        self.assertEqual(list(existente.vagas_interesse.all()), [self.vaga])

    def test_validacoes_do_formulario_publico(self):
        link = LinkCadastro.pessoal_de(self.rec)
        for campo, valor, trecho in [("cpf", "111.111.111-11", "CPF inválido"), ("telefone", "123", "telefone"),
                                     ("cep", "123", "CEP inválido"), ("nome", "Maria", "sobrenome")]:
            resposta = self.client.post(link.caminho, dados_publicos(1, **{campo: valor}))
            self.assertEqual(resposta.status_code, 200)
            self.assertContains(resposta, trecho)
        sem_aceite = dados_publicos(1)
        del sem_aceite["aceite"]
        self.assertContains(self.client.post(link.caminho, sem_aceite), "autorizar o uso dos dados")
        self.assertEqual(Candidato.objects.count(), 0)

    def test_isca_de_robo_finge_sucesso_sem_gravar(self):
        link = LinkCadastro.pessoal_de(self.rec)
        resposta = self.client.post(link.caminho, dados_publicos(1, website="http://spam"))
        self.assertRedirects(resposta, reverse("candidatos:cadastro_concluido"))
        self.assertEqual(Candidato.objects.count(), 0)

    def test_paginas_publicas_nao_vazam_o_token_pelo_referer(self):
        link = LinkCadastro.pessoal_de(self.rec)
        self.assertEqual(self.client.get(link.caminho)["Referrer-Policy"], "no-referrer")


class ImportacaoTests(BaseTests):
    def planilha(self, linhas):
        modelo = load_workbook(BytesIO(planilhas.gerar_modelo(self.vaga, Configuracao.atual())))
        aba = modelo["Candidatos"]
        for i, linha in enumerate(linhas, start=2):
            for j, valor in enumerate(linha, start=1):
                aba.cell(row=i, column=j, value=valor)
        saida = BytesIO()
        modelo.save(saida)
        return SimpleUploadedFile("candidatos.xlsx", saida.getvalue())

    def test_modelo_traz_nome_cpf_telefone_endereco_e_identifica_a_vaga(self):
        wb = load_workbook(BytesIO(planilhas.gerar_modelo(self.vaga, Configuracao.atual())))
        cabecalho = [c.value for c in wb["Candidatos"][1]]
        self.assertEqual(cabecalho[:4], ["Nome completo *", "CPF *", "Telefone (WhatsApp) *", "CEP *"])
        self.assertIn("Endereço (rua ou avenida)", cabecalho)
        self.assertEqual(wb["_rrpa"]["B2"].value, self.vaga.pk)

    def test_importacao_com_previa_e_confirmacao(self):
        existente = Candidato.objects.create(nome="Já Existe Aqui", cpf=cpf_valido(21), telefone=telefone_ficticio(21), cep="01310100")
        arquivo = self.planilha([
            ["MARIA DA SILVA", cpf_valido(20), "(11) 90000-0020", "01310-100", "", "", "", "", "", ""],      # nova (nome em maiúsculas)
            ["Já Existe Aqui", cpf_valido(21), telefone_ficticio(21), "01310100", "", "", "", "", "", ""],   # duplicada
            ["Sem Cpf Valido", "123", telefone_ficticio(22), "01310100", "", "", "", "", "", ""],           # erro
            ["Repetida Na Planilha", cpf_valido(20), telefone_ficticio(23), "01310100", "", "", "", "", "", ""],  # CPF repetido
        ])
        self.client.force_login(self.rec)
        resposta = self.client.post(reverse("candidatos:importar_vaga", args=[self.vaga.pk]), {"arquivo": arquivo})
        imp = Importacao.objects.get()
        self.assertRedirects(resposta, reverse("candidatos:previa", args=[imp.pk]))
        self.assertEqual((imp.total, imp.novas, imp.duplicadas, imp.erros), (4, 1, 1, 2))
        self.assertEqual(Candidato.objects.count(), 1)  # a prévia não cadastra ninguém
        self.assertContains(self.client.get(reverse("candidatos:previa", args=[imp.pk])), "Linhas com problema")

        erros = self.client.get(reverse("candidatos:erros_xlsx", args=[imp.pk]))
        self.assertEqual(len(list(load_workbook(BytesIO(erros.content))["Candidatos"].iter_rows())), 3)  # cabeçalho + 2 linhas

        self.client.post(reverse("candidatos:confirmar_importacao", args=[imp.pk]))
        imp.refresh_from_db()
        self.assertEqual((imp.status, imp.criadas, imp.vinculadas), ("confirmada", 1, 1))
        nova = Candidato.objects.get(cpf=cpf_valido(20))
        self.assertEqual((nova.nome, nova.origem, nova.status, nova.criado_por), ("Maria da Silva", "planilha", "pendente", self.rec))
        self.assertEqual(self.vaga.candidatos.count(), 2)  # a nova e a que já existia
        existente.refresh_from_db()
        self.assertEqual(existente.nome, "Já Existe Aqui")
        # confirmar de novo não duplica
        self.client.post(reverse("candidatos:confirmar_importacao", args=[imp.pk]))
        self.assertEqual(Candidato.objects.count(), 2)

    def test_arquivo_sem_titulos_e_recusado_com_mensagem_clara(self):
        wb = Workbook()
        wb.active.append(["a", "b"])
        saida = BytesIO()
        wb.save(saida)
        self.client.force_login(self.rec)
        resposta = self.client.post(reverse("candidatos:importar_vaga", args=[self.vaga.pk]),
                                    {"arquivo": SimpleUploadedFile("x.xlsx", saida.getvalue())})
        self.assertContains(resposta, "Não encontrei os títulos das colunas")
        self.assertEqual(Importacao.objects.count(), 0)

    def test_csv_com_ponto_e_virgula(self):
        csv = f"Nome;CPF;Telefone;CEP\nCarlos Mendes;{cpf_valido(30)};{telefone_ficticio(30)};01310100\n".encode()
        self.client.force_login(self.rec)
        self.client.post(reverse("candidatos:importar_vaga", args=[self.vaga.pk]), {"arquivo": SimpleUploadedFile("c.csv", csv)})
        self.assertEqual(Importacao.objects.get().novas, 1)

    def test_gestor_nao_importa(self):
        self.client.force_login(self.gestor)
        self.assertEqual(self.client.get(reverse("candidatos:importar_vaga", args=[self.vaga.pk])).status_code, 403)


class FluxoDeAprovacaoTests(BaseTests):
    def setUp(self):
        self.c = Candidato.objects.create(nome="Pedro Alves", cpf=cpf_valido(900), telefone=telefone_ficticio(9), cep="01310100")

    def test_aprovar_reprovar_e_reabrir(self):
        self.client.force_login(self.rec)
        self.client.post(reverse("candidatos:aprovar", args=[self.c.pk]))
        self.c.refresh_from_db()
        self.assertEqual((self.c.status, self.c.aprovado_por), ("aprovado", self.rec))
        self.client.post(reverse("candidatos:reprovar", args=[self.c.pk]), {"motivo": "Sem disponibilidade"})
        self.c.refresh_from_db()
        self.assertEqual(self.c.status, "reprovado")
        self.assertIn("Sem disponibilidade", self.c.observacoes)
        self.client.post(reverse("candidatos:reabrir", args=[self.c.pk]))
        self.c.refresh_from_db()
        self.assertEqual(self.c.status, "pendente")

    def test_gestor_nao_aprova(self):
        self.client.force_login(self.gestor)
        self.assertEqual(self.client.post(reverse("candidatos:aprovar", args=[self.c.pk])).status_code, 403)
        self.c.refresh_from_db()
        self.assertEqual(self.c.status, "pendente")

    def test_enviar_convite_aprova_e_leva_para_a_mensagem(self):
        self.client.force_login(self.rec)
        resposta = self.client.post(reverse("candidatos:convite", args=[self.c.pk]))
        link = LinkCadastro.objects.get()
        self.assertRedirects(resposta, f"{reverse('mensagens:compor')}?candidato={self.c.pk}&link={link.pk}&categoria=convite_cadastro",
                             fetch_redirect_response=False)
        self.c.refresh_from_db()
        self.assertEqual(self.c.status, "aprovado")

    def test_anonimizar_remove_dados_pessoais_e_so_o_admin_faz(self):
        self.client.force_login(self.rec)
        self.assertEqual(self.client.post(reverse("candidatos:anonimizar", args=[self.c.pk])).status_code, 403)
        admin = Usuario.objects.create_user("adm", password="x", perfil="admin")
        self.client.force_login(admin)
        self.client.post(reverse("candidatos:anonimizar", args=[self.c.pk]))
        self.c.refresh_from_db()
        self.assertEqual((self.c.nome, self.c.telefone, self.c.cep, self.c.status), ("Cadastro removido", "", "", "desistente"))

    def test_registrar_que_pessoa_nao_quer_whatsapp(self):
        self.client.force_login(self.rec)
        self.client.post(reverse("candidatos:whatsapp", args=[self.c.pk]), {"valor": "nao"})
        self.c.refresh_from_db()
        self.assertIs(self.c.whatsapp_ok, False)


class ExclusaoDeCandidatoTests(BaseTests):
    def setUp(self):
        self.admin = Usuario.objects.create_user("adm", password="x", perfil="admin")
        self.c = Candidato.objects.create(nome="Pedro Alves", cpf=cpf_valido(900), telefone=telefone_ficticio(9), cep="01310100")

    def outro(self, n, **extra):
        return Candidato.objects.create(nome=f"Pessoa {n}", cpf=cpf_valido(900 + n), telefone=telefone_ficticio(n), **extra)

    def chamado(self, n=1):
        """Candidato que já foi chamado para uma vaga (e o vínculo criado)."""
        candidato = self.outro(n, status=Candidato.Status.APROVADO)
        return candidato, vinculos.chamar_para_vaga(candidato, self.vaga, self.rec)

    def excluir(self, candidato, usuario=None, **dados):
        self.client.force_login(usuario or self.admin)
        return self.client.post(reverse("candidatos:excluir", args=[candidato.pk]), dados)

    def test_admin_exclui_o_cadastro_e_volta_para_a_lista(self):
        resposta = self.excluir(self.c)
        self.assertRedirects(resposta, reverse("candidatos:lista"))
        self.assertFalse(Candidato.objects.filter(pk=self.c.pk).exists())
        log = LogAtividade.objects.get(acao="candidato.excluido")
        self.assertEqual((log.usuario, log.descricao, log.entidade_id), (self.admin, "Excluiu o cadastro de Pedro Alves", self.c.pk))
        self.assertIn("O cadastro de Pedro Alves foi excluído.", [str(m) for m in get_messages(resposta.wsgi_request)])

    def test_quem_esta_em_vaga_e_excluido_com_o_vinculo_e_os_checkins_e_a_vaga_fica_livre(self):
        candidato, vinculo = self.chamado(1)
        Checkin.objects.create(vinculo=vinculo, competencia=vinculo.inicio_periodo.replace(day=1),
                               resultado=Checkin.Resultado.CUMPRIU, registrado_por=self.rec)
        _, vinculo_de_outro = self.chamado(2)
        self.assertEqual(Vinculo.objects.ocupando().filter(vaga=self.vaga).count(), 2)

        self.excluir(candidato)

        self.assertFalse(Candidato.objects.filter(pk=candidato.pk).exists())
        self.assertFalse(Vinculo.objects.filter(pk=vinculo.pk).exists())
        self.assertEqual(Checkin.objects.count(), 0)
        self.assertEqual(list(Vinculo.objects.ocupando().filter(vaga=self.vaga)), [vinculo_de_outro])  # a vaga ficou livre; o outro segue
        log = LogAtividade.objects.get(acao="candidato.excluido")
        self.assertEqual(log.descricao, "Excluiu o cadastro de Pessoa 1 e 1 vínculo(s) com 1 check-in(s)")

    def test_vinculo_ja_encerrado_tambem_e_apagado(self):
        candidato, vinculo = self.chamado()
        vinculos.encerrar_vinculo(vinculo, self.rec, "desistencia")
        self.excluir(candidato)
        self.assertFalse(Candidato.objects.filter(pk=candidato.pk).exists())
        self.assertFalse(Vinculo.objects.exists())

    def test_mensagens_e_convites_ficam_sem_dados_pessoais(self):
        disparo = Disparo.objects.create(nome="Convocação")
        mensagem = MensagemDisparo.objects.create(disparo=disparo, candidato=self.c, nome_destino="Pedro Alves", telefone="11999990000",
                                                  texto="Oi, Pedro", status=MensagemDisparo.Status.FILA)
        resposta = MensagemRecebida.objects.create(candidato=self.c, telefone="5511999990000", texto="SAIR")
        envio = EnvioMensagem.objects.create(candidato=self.c, nome_destino="Pedro Alves", telefone="11999990000", texto="Oi, Pedro")
        convite = services.criar_convite(self.rec, "Pedro Alves", "11999990000", candidato=self.c)

        self.excluir(self.c)

        for objeto in (mensagem, resposta, envio, convite):
            objeto.refresh_from_db()
        self.assertEqual((mensagem.candidato, mensagem.telefone, mensagem.texto, mensagem.status), (None, "", "", "ignorada"))
        self.assertEqual((resposta.candidato, resposta.telefone, resposta.texto), (None, "", ""))
        self.assertEqual((envio.candidato, envio.nome_destino, envio.telefone, envio.texto), (None, "Cadastro removido", "", ""))
        self.assertEqual((convite.nome_sugerido, convite.telefone_sugerido), ("", ""))
        self.assertEqual(convite.estado, "revogado")  # o link aberto não pode mais ser usado

    def test_convite_ja_usado_nao_e_marcado_como_revogado(self):
        convite = services.criar_convite(self.rec, "Pedro Alves", "11999990000", candidato=self.c)
        convite.usado_em = timezone.now()
        convite.save()
        self.excluir(self.c)
        convite.refresh_from_db()
        self.assertIsNone(convite.revogado_em)
        self.assertEqual(convite.telefone_sugerido, "")

    def test_se_algo_protege_o_cadastro_nada_e_alterado(self):
        candidato, vinculo = self.chamado()
        envio = EnvioMensagem.objects.create(candidato=candidato, nome_destino="Pessoa 1", telefone="11999990000", texto="Oi")
        with mock.patch.object(Candidato, "delete", side_effect=ProtectedError("protegido", set())):
            with self.assertRaises(RegraDeNegocio):
                services.excluir(candidato, self.admin)
        envio.refresh_from_db()
        self.assertEqual((envio.candidato, envio.texto), (candidato, "Oi"))
        self.assertTrue(Vinculo.objects.filter(pk=vinculo.pk).exists())  # o vínculo apagado antes foi desfeito
        self.assertFalse(LogAtividade.objects.filter(acao="candidato.excluido").exists())

    def test_so_o_admin_exclui(self):
        for usuario in (self.rec, self.gestor):
            self.assertEqual(self.excluir(self.c, usuario).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.post(reverse("candidatos:excluir", args=[self.c.pk])).status_code, 302)  # vai para a entrada
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("candidatos:excluir", args=[self.c.pk])).status_code, 405)  # só POST
        self.assertTrue(Candidato.objects.filter(pk=self.c.pk).exists())

    def test_candidato_inexistente_da_404(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.post(reverse("candidatos:excluir", args=[999999])).status_code, 404)

    def test_exclusao_em_massa_apaga_os_selecionados_e_so_eles(self):
        livre = self.outro(1)
        em_vaga, _ = self.chamado(2)
        fica = self.outro(3)
        self.client.force_login(self.admin)
        resposta = self.client.post(reverse("candidatos:acoes"), {"acao": "excluir", "ids": [self.c.pk, livre.pk, em_vaga.pk]})
        self.assertRedirects(resposta, reverse("candidatos:lista"))
        self.assertEqual(list(Candidato.objects.values_list("pk", flat=True)), [fica.pk])
        self.assertFalse(Vinculo.objects.exists())
        self.assertEqual([str(m) for m in get_messages(resposta.wsgi_request)], ["3 candidato(s) excluído(s)."])

    def test_exclusao_em_massa_e_so_do_admin_mas_as_outras_acoes_continuam(self):
        self.client.force_login(self.rec)
        resposta = self.client.post(reverse("candidatos:acoes"), {"acao": "excluir", "ids": [self.c.pk]})
        self.assertEqual(resposta.status_code, 403)
        self.assertTrue(Candidato.objects.filter(pk=self.c.pk).exists())
        self.client.post(reverse("candidatos:acoes"), {"acao": "aprovar", "ids": [self.c.pk]})
        self.c.refresh_from_db()
        self.assertEqual(self.c.status, "aprovado")

    def test_botoes_so_aparecem_para_o_admin_e_a_confirmacao_avisa_do_vinculo(self):
        url_excluir = reverse("candidatos:excluir", args=[self.c.pk])
        self.client.force_login(self.admin)
        ficha = self.client.get(self.c.get_absolute_url())
        self.assertContains(ficha, f'action="{url_excluir}"')
        self.assertContains(ficha, 'data-confirmar="Excluir o cadastro de Pedro Alves? Não pode ser desfeito."')
        self.assertContains(self.client.get(reverse("candidatos:lista")), 'value="excluir"')

        chamado, vinculo = self.chamado()
        Checkin.objects.create(vinculo=vinculo, competencia=vinculo.inicio_periodo.replace(day=1), resultado=Checkin.Resultado.CUMPRIU)
        ficha = self.client.get(chamado.get_absolute_url())
        self.assertContains(ficha, "Isso apaga também 1 vínculo(s) e 1 check-in(s) e libera a vaga.")

        self.client.force_login(self.rec)
        self.assertNotContains(self.client.get(self.c.get_absolute_url()), url_excluir)
        self.assertNotContains(self.client.get(reverse("candidatos:lista")), 'value="excluir"')
