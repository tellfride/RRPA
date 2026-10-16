"""Leitura, validação e importação das planilhas de candidatos (.xlsx ou .csv)."""
import csv
import io
import re
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from core import geo
from core.models import LogAtividade
from core.utils import (
    limpar_nome, normalizar_cep, normalizar_cpf, normalizar_telefone, sem_acentos, texto_de_planilha, validar_cep,
    validar_cpf, validar_telefone,
)

from .forms import UFS
from .models import Candidato, Importacao

LIMITE_LINHAS = 5000
SINONIMOS = {
    "nome": {"nome", "nome completo", "candidato", "colaborador"},
    "cpf": {"cpf"},
    "telefone": {"telefone", "celular", "whatsapp", "fone", "telefone whatsapp", "telefone com whatsapp"},
    "cep": {"cep"},
    "logradouro": {"endereco", "logradouro", "rua", "endereco rua ou avenida", "rua ou avenida"},
    "numero": {"numero", "n", "num", "no"},
    "complemento": {"complemento"},
    "bairro": {"bairro"},
    "cidade": {"cidade", "municipio"},
    "uf": {"uf", "estado"},
}
ROTULOS = {"nome": "Nome completo", "cpf": "CPF", "telefone": "Telefone (WhatsApp)", "cep": "CEP", "logradouro": "Endereço",
           "numero": "Número", "complemento": "Complemento", "bairro": "Bairro", "cidade": "Cidade", "uf": "UF"}


class ErroPlanilha(Exception):
    """Problema que impede de ler o arquivo; a mensagem é mostrada ao usuário."""


def _normalizar_titulo(valor):
    texto = sem_acentos(str(valor or "")).lower()
    return " ".join(re.sub(r"[^a-z0-9 ]", " ", texto).split())


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _mapear_cabecalho(linha):
    mapa = {}
    for indice, valor in enumerate(linha):
        titulo = _normalizar_titulo(valor)
        for campo, nomes in SINONIMOS.items():
            if titulo in nomes and campo not in mapa:
                mapa[campo] = indice
    return mapa


def _linhas_xlsx(arquivo):
    try:
        wb = load_workbook(arquivo, read_only=True, data_only=True)
    except Exception as exc:
        raise ErroPlanilha("Não consegui abrir este arquivo. Use uma planilha .xlsx (o modelo da vaga) ou um .csv.") from exc
    nomes = [n for n in wb.sheetnames if n not in ("_rrpa", "Instruções") and wb[n].sheet_state == "visible"]
    if not nomes:
        raise ErroPlanilha("A planilha não tem nenhuma aba com dados.")
    aba = "Candidatos" if "Candidatos" in nomes else nomes[0]
    return [list(linha) for linha in wb[aba].iter_rows(values_only=True)]


def _linhas_csv(arquivo):
    bruto = arquivo.read()
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = bruto.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErroPlanilha("Não consegui ler o texto do arquivo. Salve o CSV em UTF-8.")
    amostra = texto[:2048]
    delimitador = ";" if amostra.count(";") >= amostra.count(",") else ","
    return [linha for linha in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def ler_planilha(arquivo):
    """Devolve a lista de linhas (dicts campo -> texto, com o número da linha original)."""
    nome = (arquivo.name or "").lower()
    if nome.endswith((".xlsx", ".xlsm")):
        todas = _linhas_xlsx(arquivo)
    elif nome.endswith(".csv"):
        todas = _linhas_csv(arquivo)
    else:
        raise ErroPlanilha("Envie um arquivo .xlsx (modelo da vaga) ou .csv.")

    indice_cabecalho, mapa = None, {}
    for i, linha in enumerate(todas[:15]):
        candidato = _mapear_cabecalho(linha)
        if "nome" in candidato and "cpf" in candidato:
            indice_cabecalho, mapa = i, candidato
            break
    if indice_cabecalho is None:
        raise ErroPlanilha("Não encontrei os títulos das colunas (Nome, CPF, Telefone e CEP). Use o modelo de planilha da vaga.")
    faltando = [ROTULOS[c] for c in ("telefone", "cep") if c not in mapa]
    if faltando:
        raise ErroPlanilha("Faltam colunas na planilha: " + ", ".join(faltando) + ". Use o modelo de planilha da vaga.")

    linhas = []
    for numero, bruta in enumerate(todas[indice_cabecalho + 1:], start=indice_cabecalho + 2):
        dados = {campo: _texto(bruta[i]) if i < len(bruta) else "" for campo, i in mapa.items()}
        if not any(dados.values()):
            continue
        dados["linha"] = numero
        linhas.append(dados)
        if len(linhas) > LIMITE_LINHAS:
            raise ErroPlanilha(f"A planilha tem mais de {LIMITE_LINHAS} linhas. Divida em arquivos menores.")
    if not linhas:
        raise ErroPlanilha("A planilha está vazia. Preencha ao menos uma linha abaixo dos títulos.")
    return linhas


def validar_linhas(linhas):
    """Marca cada linha como nova, duplicada (CPF já cadastrado) ou erro, com os motivos."""
    cpfs = {normalizar_cpf(l.get("cpf")) for l in linhas}
    existentes = {c.cpf: c.nome for c in Candidato.objects.filter(cpf__in=cpfs).only("cpf", "nome")}
    vistos = {}
    resultado = []
    for dados in linhas:
        erros = []
        nome = limpar_nome(dados.get("nome"))
        cpf = normalizar_cpf(dados.get("cpf"))
        telefone = normalizar_telefone(dados.get("telefone"))
        cep = normalizar_cep(dados.get("cep"))
        uf = (dados.get("uf") or "").strip().upper()
        if len(nome.split()) < 2:
            erros.append("Informe nome e sobrenome.")
        if not validar_cpf(cpf):
            erros.append("CPF inválido.")
        elif cpf in vistos:
            erros.append(f"CPF repetido na planilha (linha {vistos[cpf]}).")
        if not validar_telefone(telefone):
            erros.append("Telefone inválido: use DDD + número.")
        if not validar_cep(cep):
            erros.append("CEP inválido: são 8 números.")
        if uf and uf not in UFS:
            erros.append("UF inválida.")
        if validar_cpf(cpf):
            vistos.setdefault(cpf, dados["linha"])
        estado = "erro" if erros else ("duplicada" if cpf in existentes else "nova")
        resultado.append({
            "linha": dados["linha"], "nome": nome, "cpf": cpf, "telefone": telefone, "cep": cep,
            "logradouro": dados.get("logradouro", ""), "numero": dados.get("numero", ""),
            "complemento": dados.get("complemento", ""), "bairro": dados.get("bairro", ""),
            "cidade": dados.get("cidade", ""), "uf": uf, "estado": estado, "erros": erros,
            "existente": existentes.get(cpf, ""),
        })
    return resultado


def criar_previa(arquivo, vaga, usuario, como_aprovados=False):
    linhas = validar_linhas(ler_planilha(arquivo))
    contagem = {e: sum(1 for l in linhas if l["estado"] == e) for e in ("nova", "duplicada", "erro")}
    return Importacao.objects.create(
        vaga=vaga, usuario=usuario, nome_arquivo=arquivo.name[:255], como_aprovados=como_aprovados, dados=linhas,
        total=len(linhas), novas=contagem["nova"], duplicadas=contagem["duplicada"], erros=contagem["erro"],
    )


@transaction.atomic
def confirmar(importacao, usuario, request=None):
    importacao = Importacao.objects.select_for_update().get(pk=importacao.pk)
    if importacao.status != Importacao.Status.PREVIA:
        raise ErroPlanilha("Esta importação já foi concluída ou cancelada.")
    agora = timezone.now()
    criadas = vinculadas = 0
    for linha in importacao.dados:
        if linha["estado"] == "erro":
            continue
        candidato = Candidato.objects.filter(cpf=linha["cpf"]).first()
        if candidato is None:
            candidato = Candidato(
                nome=linha["nome"], cpf=linha["cpf"], telefone=linha["telefone"], cep=linha["cep"],
                logradouro=linha["logradouro"], numero=linha["numero"], complemento=linha["complemento"],
                bairro=linha["bairro"], cidade=linha["cidade"], uf=linha["uf"], origem=Candidato.Origem.PLANILHA,
                status=Candidato.Status.APROVADO if importacao.como_aprovados else Candidato.Status.PENDENTE,
                aprovado_por=usuario if importacao.como_aprovados else None,
                aprovado_em=agora if importacao.como_aprovados else None, criado_por=usuario,
            )
            candidato.save()
            geo.aplicar_localizacao(candidato, somente_cache=True)
            criadas += 1
        else:
            vinculadas += 1
        if importacao.vaga_id:
            candidato.vagas_interesse.add(importacao.vaga)
    importacao.status = Importacao.Status.CONFIRMADA
    importacao.criadas, importacao.vinculadas, importacao.confirmada_em = criadas, vinculadas, agora
    importacao.save()
    destino = f" para {importacao.vaga.nome}" if importacao.vaga_id else ""
    LogAtividade.registrar(
        usuario, "candidato.importado",
        f"Importou {criadas} candidatos{destino} da planilha {importacao.nome_arquivo}", importacao, request,
    )
    return importacao


def planilha_de_erros(importacao):
    """Linhas com problema, no mesmo formato do modelo + coluna 'Problema', para corrigir e reenviar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidatos"
    campos = ["nome", "cpf", "telefone", "cep", "logradouro", "numero", "complemento", "bairro", "cidade", "uf"]
    ws.append([ROTULOS[c] for c in campos] + ["Problema", "Linha original"])
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0E6B66")
    for linha in importacao.dados:
        if linha["estado"] == "erro":
            ws.append([texto_de_planilha(linha[c]) for c in campos] + [" ".join(linha["erros"]), linha["linha"]])
    for coluna, largura in zip("ABCDEFGHIJKL", [34, 17, 20, 12, 36, 10, 18, 22, 22, 6, 50, 14]):
        ws.column_dimensions[coluna].width = largura
    for linha in ws.iter_rows(min_row=2, max_col=10):
        for celula in linha:
            celula.number_format = "@"
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
